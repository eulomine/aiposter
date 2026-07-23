"""
poster_engine.py
포스터 생성 엔진 — 사이즈 관리, Gemini API 호출, PPTX 조립
모든 작업을 메모리에서 처리 (디스크 저장 없음)
v5: 듀얼 모드(Puter.js / API키), 자동 폴백, 부제 지원,
    절대 규칙, 배경 프롬프트 개선, 3슬라이드, PPTX 크기 제한
"""

import base64
import math
import time
from io import BytesIO
from google import genai
from google.genai import types
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

# ============================================================
# 1. 사이즈 프리셋 & 비율 매칭
# ============================================================

GEMINI_RATIOS = [
    (1, 1), (3, 2), (2, 3), (3, 4), (4, 3),
    (4, 5), (5, 4), (9, 16), (16, 9), (21, 9),
    (9, 21), (1, 4), (4, 1), (1, 8), (8, 1),
]

SIZE_PRESETS = {
    "A3_landscape":   {"name": "A3 가로",       "width_cm": 42.0,   "height_cm": 29.7},
    "A3_portrait":    {"name": "A3 세로",       "width_cm": 29.7,   "height_cm": 42.0},
    "16:9_landscape": {"name": "16:9 가로",     "width_cm": 50.8,   "height_cm": 28.575},
    "16:9_portrait":  {"name": "16:9 세로",     "width_cm": 28.575, "height_cm": 50.8},
    "square":         {"name": "정방형 (1:1)",   "width_cm": 40.0,   "height_cm": 40.0},
    "banner_h":       {"name": "가로 배너",      "width_cm": 120.0,  "height_cm": 30.0},
    "banner_v":       {"name": "세로 배너",      "width_cm": 30.0,   "height_cm": 120.0},
    "hanging_banner": {"name": "가로 현수막",    "width_cm": 500.0,  "height_cm": 90.0},
}

# ============================================================
# 2. 스타일 프리셋
# ============================================================

STYLE_PRESETS = {
    "cute_illust": {
        "name": "귀여운 일러스트",
        "prompt_en": "cute hand-drawn illustration style, warm pastel colors, friendly cartoon characters, soft rounded shapes, children's book aesthetic",
    },
    "corporate": {
        "name": "깔끔한 기업형",
        "prompt_en": "clean modern corporate design, professional layout, geometric shapes, business-appropriate color palette, sleek typography areas, polished and minimal",
    },
    "photo_real": {
        "name": "실사 사진풍",
        "prompt_en": "photorealistic stock photography style, real-life scenes with natural lighting, professional DSLR photo quality, realistic people and environments, cinematic composition, NOT illustration NOT cartoon NOT drawing",
    },
    "watercolor": {
        "name": "수채화풍",
        "prompt_en": "beautiful watercolor painting style, soft color bleeding, artistic brush strokes, elegant and artistic mood, traditional art feeling",
    },
    "minimal": {
        "name": "미니멀 / 타이포",
        "prompt_en": "minimalist design, clean white space, bold typography focus, simple geometric accents, modern and sophisticated, limited color palette",
    },
    "retro": {
        "name": "레트로 / 빈티지",
        "prompt_en": "vintage retro design style, aged paper texture, classic typography, muted warm tones, nostalgic 1970s-80s aesthetic, worn and weathered look",
    },
    "pop_art": {
        "name": "팝아트 / 컬러풀",
        "prompt_en": "vibrant pop art style, bold bright colors, dynamic composition, energetic and eye-catching, comic book influenced, high contrast",
    },
    "elegant": {
        "name": "고급스러운 / 포멀",
        "prompt_en": "elegant formal design, dark sophisticated color scheme, gold or silver accents, luxurious feel, serif typography areas, premium quality aesthetic",
    },
    "custom": {
        "name": "직접 입력",
        "prompt_en": "",
    },
}


def find_closest_gemini_ratio(width_cm: float, height_cm: float) -> str:
    target = width_cm / height_cm
    best_ratio = None
    best_diff = float('inf')
    for w, h in GEMINI_RATIOS:
        diff = abs(math.log(target) - math.log(w / h))
        if diff < best_diff:
            best_diff = diff
            best_ratio = (w, h)
    return f"{best_ratio[0]}:{best_ratio[1]}"


def get_size_info(preset_key=None, custom_w=None, custom_h=None):
    if preset_key and preset_key in SIZE_PRESETS:
        info = SIZE_PRESETS[preset_key].copy()
        info["gemini_ratio"] = find_closest_gemini_ratio(info["width_cm"], info["height_cm"])
        return info
    elif custom_w and custom_h:
        return {
            "name": f"커스텀 ({custom_w}×{custom_h}cm)",
            "width_cm": float(custom_w),
            "height_cm": float(custom_h),
            "gemini_ratio": find_closest_gemini_ratio(float(custom_w), float(custom_h)),
        }
    raise ValueError("사이즈를 지정해주세요.")


# ============================================================
# 3. Gemini API 호출 (자동 폴백: 고화질 → 무료모델)
# ============================================================

def init_gemini(api_key: str):
    return genai.Client(api_key=api_key)


def validate_api_key(client):
    try:
        client.models.generate_content(
            model="gemini-3.5-flash",
            contents="Say OK",
            config=types.GenerateContentConfig(max_output_tokens=5),
        )
        return True
    except Exception as e:
        if "API key not valid" in str(e):
            return False
        return True


def generate_image(client, prompt, aspect_ratio, image_size="2K",
                   ref_image_bytes=None,
                   model="gemini-3.1-flash-image",
                   fallback_model="gemini-3.1-flash-lite-image",
                   max_retries=2):
    """
    이미지 생성 — 자동 폴백 지원.
    1) model (고화질)로 시도
    2) 실패 시 fallback_model (무료/저화질)로 재시도
    반환: (image_bytes, used_model_name)
    """
    models_to_try = [
        (model, image_size),
        (fallback_model, "1K"),  # lite 모델은 1K만 지원
    ]

    last_error = None
    for current_model, current_size in models_to_try:
        for attempt in range(max_retries + 1):
            try:
                contents = []
                if ref_image_bytes:
                    contents.append(
                        types.Part.from_bytes(data=ref_image_bytes, mime_type="image/jpeg")
                    )
                contents.append(types.Part.from_text(text=prompt))

                response = client.models.generate_content(
                    model=current_model,
                    contents=contents,
                    config=types.GenerateContentConfig(
                        response_modalities=["IMAGE", "TEXT"],
                    ),
                )

                for part in response.candidates[0].content.parts:
                    if part.inline_data and part.inline_data.mime_type.startswith("image/"):
                        return part.inline_data.data, current_model

                raise RuntimeError("Gemini가 이미지를 반환하지 않았습니다.")

            except Exception as e:
                last_error = e
                error_str = str(e)
                # 즉시 폴백해야 하는 에러 (quota, 모델 사용 불가 등)
                if any(keyword in error_str.lower() for keyword in
                       ["quota", "429", "resource exhausted", "not available",
                        "not found", "permission"]):
                    break  # 다음 모델로 넘어감
                # 일시적 에러 → 재시도
                if attempt < max_retries:
                    time.sleep(5 * (attempt + 1))
                    continue
                break  # 재시도 다 소진, 다음 모델로

    raise last_error


def generate_text_plan(client, poster_info, model="gemini-3.5-flash"):
    """텍스트 모델로 디자인 계획 생성 — 절대 규칙 포함"""

    upload_desc = poster_info.get('upload_descriptions', '없음')
    has_uploads = upload_desc != '없음' and upload_desc.strip()

    if has_uploads:
        upload_instruction = f"""업로드된 이미지: {upload_desc}
- 각 이미지의 설명을 참고하여 적절한 위치에 배치를 계획하세요."""
    else:
        upload_instruction = """업로드된 이미지: 없음
- 로고, 사진, QR코드 등의 공간을 미리 확보하지 마세요.
- 전체를 그래픽과 텍스트로만 구성하세요."""

    info_items = []
    for key, label in [
        ('title', '제목'), ('subtitle', '부제'), ('date', '날짜'),
        ('venue', '장소'), ('organizer', '주관'), ('content', '내용'),
        ('additional', 'AI 추가 표현 요청'),
    ]:
        val = poster_info.get(key, '')
        if val and val.strip():
            info_items.append(f"- {label}: {val}")

    info_block = "\n".join(info_items) if info_items else "- (입력된 정보 없음)"

    prompt = f"""당신은 전문 포스터 디자이너입니다. 다음 정보로 포스터 디자인을 계획해주세요.

[절대 규칙 — 반드시 지키세요]
1. 아래에 직접 입력된 정보만 사용하세요. 전화번호, 이메일, 홈페이지, 등록방법, 문의처, 담당자 이름 등을 추측하거나 임의로 만들어내지 마세요.
2. 아래에 없는 항목은 디자인에서 완전히 제외하세요.
3. 업로드된 이미지가 없으면 로고/사진/QR코드 공간을 만들지 마세요.

포스터 정보:
{info_block}
- 분위기/스타일: {poster_info.get('mood', '')}
- 크기: {poster_info.get('size_name', '')} (비율: {poster_info.get('gemini_ratio', '')})
{upload_instruction}

다음을 포함해서 답변해주세요:
1. 전체 분위기와 색감 설명
2. 배경 이미지 생성용 프롬프트 (영어, 상세하게)
3. 텍스트 배치 계획 (위치, 크기, 색상) — 입력된 항목만
4. 업로드 이미지가 있을 경우에만 이미지 배치 제안
5. 추천 장식 요소"""

    response = client.models.generate_content(model=model, contents=prompt)
    return response.text


# ============================================================
# 4. PPTX 생성 — 3슬라이드 구조
# ============================================================

def create_poster_pptx(
    width_cm, height_cm,
    background_bytes,
    original_bytes,
    reference_bytes,
    texts,
    upload_images=None,
    title="poster",
    actual_width_cm=None,
    actual_height_cm=None,
):
    prs = Presentation()

    MAX_CM = 142.0
    pptx_scale = 1.0
    if width_cm > MAX_CM or height_cm > MAX_CM:
        pptx_scale = min(MAX_CM / width_cm, MAX_CM / height_cm)
        width_cm = round(width_cm * pptx_scale, 2)
        height_cm = round(height_cm * pptx_scale, 2)

    prs.slide_width = Cm(width_cm)
    prs.slide_height = Cm(height_cm)

    blank_layout = prs.slide_layouts[6]

    # ── 슬라이드 1: 편집용 포스터 ──
    slide1 = prs.slides.add_slide(blank_layout)
    slide1.shapes.add_picture(
        BytesIO(background_bytes), Cm(0), Cm(0),
        width=Cm(width_cm), height=Cm(height_cm)
    )

    align_map = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }

    for t in texts:
        txBox = slide1.shapes.add_textbox(
            Cm(t["x_cm"]), Cm(t["y_cm"]),
            Cm(t["width_cm"]), Cm(t["height_cm"]),
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = t["content"]
        p.font.size = Pt(t.get("font_size_pt", 24))
        p.font.name = t.get("font_name", "맑은 고딕")
        p.font.bold = t.get("bold", False)
        color_hex = t.get("font_color", "FFFFFF")
        p.font.color.rgb = RGBColor.from_string(color_hex)
        p.alignment = align_map.get(t.get("align", "center"), PP_ALIGN.CENTER)

    if upload_images:
        for img in upload_images:
            slide1.shapes.add_picture(
                BytesIO(img["image_bytes"]),
                Cm(img["x_cm"]), Cm(img["y_cm"]),
                width=Cm(img["width_cm"]), height=Cm(img["height_cm"]),
            )

    if actual_width_cm and actual_height_cm and pptx_scale < 1.0:
        note_w = min(width_cm - 2, 30)
        size_note = slide1.shapes.add_textbox(Cm(1), Cm(height_cm - 4), Cm(note_w), Cm(3))
        sp = size_note.text_frame.paragraphs[0]
        sp.text = f"📐 실제 인쇄 크기: {actual_width_cm} × {actual_height_cm} cm (PPTX 최대 제한으로 축소됨, 비율은 동일)"
        sp.font.size = Pt(10)
        sp.font.color.rgb = RGBColor(255, 200, 0)

    # ── 슬라이드 2: 원본 활용용 ──
    slide2 = prs.slides.add_slide(blank_layout)
    slide2.shapes.add_picture(
        BytesIO(original_bytes), Cm(0), Cm(0),
        width=Cm(width_cm), height=Cm(height_cm)
    )
    note_w = min(width_cm - 2, 25)
    note_box = slide2.shapes.add_textbox(Cm(1), Cm(1), Cm(note_w), Cm(3))
    note_p = note_box.text_frame.paragraphs[0]
    note_p.text = "📌 AI 원본 이미지입니다. 텍스트가 자연스러운 부분은 그대로 활용하고, 오타 부분만 텍스트 박스로 덮어 수정하세요."
    note_p.font.size = Pt(12)
    note_p.font.color.rgb = RGBColor(255, 165, 0)

    # ── 슬라이드 3: 참고용 ──
    slide3 = prs.slides.add_slide(blank_layout)
    slide3.shapes.add_picture(
        BytesIO(reference_bytes), Cm(0), Cm(0),
        width=Cm(width_cm), height=Cm(height_cm)
    )
    note_box3 = slide3.shapes.add_textbox(Cm(1), Cm(1), Cm(note_w), Cm(3))
    note_p3 = note_box3.text_frame.paragraphs[0]
    note_p3.text = "📌 참고용 이미지입니다. 편집 완료 후 이 슬라이드를 삭제하세요."
    note_p3.font.size = Pt(12)
    note_p3.font.color.rgb = RGBColor(255, 0, 0)

    output = BytesIO()
    prs.save(output)
    output.seek(0)
    return output
